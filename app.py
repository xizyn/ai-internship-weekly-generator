import streamlit as st
from datetime import date, timedelta, datetime
from io import BytesIO
import time

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# =====================
# 页面基础设置
# =====================

st.set_page_config(
    page_title="AI实习周记生成助手",
    page_icon="📝",
    layout="centered"
)

st.title("AI实习周记生成助手 📝")

st.info(
    "输入实习单位、岗位、实习时间和休息制度，系统会自动计算实习周数、休息时间、到岗天数和工时，并导出 Word 表格版周记。"
)

st.markdown("---")


# =====================
# 基本信息输入
# =====================

unit = st.text_input(
    "实习单位",
    placeholder="请输入实习单位名称"
)

job = st.text_input(
    "实习岗位",
    placeholder="请输入实习岗位"
)
col1, col2 = st.columns(2)

with col1:
    start_date_text = st.text_input(
        "实习开始日期",
        value="2026-01-04",
        help="请按格式输入：2026-01-04"
    )

with col2:
    end_date_text = st.text_input(
        "实习结束日期",
        value="2026-04-04",
        help="请按格式输入：2026-04-04"
    )


def parse_single_date(text, field_name):
    try:
        return datetime.strptime(text.strip(), "%Y-%m-%d").date()
    except Exception:
        st.error(f"{field_name}格式错误，请使用 2026-01-04 这种格式。")
        st.stop()


start_date = parse_single_date(start_date_text, "实习开始日期")
end_date = parse_single_date(end_date_text, "实习结束日期")

if end_date < start_date:
    st.error("实习结束日期不能早于实习开始日期，请重新填写。")
    st.stop()


work_hours_per_day = st.number_input(
    "每天工作小时数",
    min_value=1.0,
    max_value=12.0,
    value=6.0,
    step=0.5
)


# =====================
# 休息制度设置
# =====================

weekday_options = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]

weekday_map = {
    "周一": 0,
    "周二": 1,
    "周三": 2,
    "周四": 3,
    "周五": 4,
    "周六": 5,
    "周日": 6,
}

rest_mode = st.selectbox(
    "休息制度",
    ["双休", "单休", "大小周", "排休/自定义日期", "无固定休息"]
)

fixed_rest_weekdays = set()
big_week_rest = set()
small_week_rest = set()
first_week_big = False

if rest_mode == "双休":
    rest_days = st.multiselect(
        "每周固定休息日",
        weekday_options,
        default=["周六", "周日"]
    )
    fixed_rest_weekdays = {weekday_map[d] for d in rest_days}

elif rest_mode == "单休":
    rest_day = st.selectbox(
        "每周固定休息日",
        weekday_options,
        index=6
    )
    fixed_rest_weekdays = {weekday_map[rest_day]}

elif rest_mode == "大小周":
    st.write("大小周说明：大周通常休1天，小周通常休2天，系统会按周交替计算。")

    first_week_big = st.radio(
        "第一周属于",
        ["大周（休1天）", "小周（休2天）"],
        horizontal=True
    ) == "大周（休1天）"

    big_rest_day = st.selectbox(
        "大周休息日",
        weekday_options,
        index=6
    )

    small_rest_days = st.multiselect(
        "小周休息日",
        weekday_options,
        default=["周六", "周日"]
    )

    big_week_rest = {weekday_map[big_rest_day]}
    small_week_rest = {weekday_map[d] for d in small_rest_days}

elif rest_mode == "排休/自定义日期":
    st.write("排休模式下，系统不会自动按星期计算休息，需要你在下方输入具体休息日期。")

elif rest_mode == "无固定休息":
    st.write("无固定休息模式下，系统不会自动按星期排休，只会识别你填写的排休日期或节假日。")


# =====================
# 自定义休息和节假日
# =====================

custom_rest_input = st.text_area(
    "排休 / 自定义休息日期（可选）",
    placeholder=(
        "一行写一个日期或时间段，例如：\n"
        "2026-01-10\n"
        "2026-01-18\n"
        "2026-02-15~2026-02-23"
    ),
    height=100
)

holiday_input = st.text_area(
    "法定节假日 / 额外休息日期（可选）",
    placeholder=(
        "一行写一个日期或时间段，例如：\n"
        "2026-02-15~2026-02-23\n"
        "2026-04-04"
    ),
    height=100
)


# =====================
# 每周工作内容
# =====================

weekly_input = st.text_area(
    "每周工作内容（可选，一行对应一周）",
    placeholder=(
        "例如：\n"
        "熟悉单位环境，了解岗位流程。\n"
        "协助完成日常事务，学习工作规范。\n"
        "参与资料整理，配合工作人员完成基础工作。"
    ),
    height=160
)

style = st.selectbox(
    "生成风格",
    ["通用正式版", "学生实习版", "简洁记录版"]
)


# =====================
# 日期工具函数
# =====================

def format_date(d):
    return f"{d.year}年{d.month}月{d.day}日"


def parse_date(text):
    return datetime.strptime(text.strip(), "%Y-%m-%d").date()


def parse_date_lines(text):
    dates = set()
    errors = []

    for line in text.splitlines():
        line = line.strip()

        if not line:
            continue

        line = line.replace("—", "~").replace("至", "~").replace("到", "~")

        try:
            if "~" in line:
                left, right = line.split("~", 1)
                d1 = parse_date(left)
                d2 = parse_date(right)

                if d2 < d1:
                    d1, d2 = d2, d1

                current = d1
                while current <= d2:
                    dates.add(current)
                    current += timedelta(days=1)
            else:
                dates.add(parse_date(line))
        except Exception:
            errors.append(line)

    return dates, errors


def format_date_ranges(dates):
    if not dates:
        return "无"

    dates = sorted(dates)
    ranges = []

    start = dates[0]
    prev = dates[0]

    for d in dates[1:]:
        if d == prev + timedelta(days=1):
            prev = d
        else:
            if start == prev:
                ranges.append(format_date(start))
            else:
                ranges.append(f"{format_date(start)}—{format_date(prev)}")
            start = d
            prev = d

    if start == prev:
        ranges.append(format_date(start))
    else:
        ranges.append(f"{format_date(start)}—{format_date(prev)}")

    return "；".join(ranges)


# =====================
# 自动计算总周数和休息日期
# =====================

total_days = (end_date - start_date).days + 1
week_count = (total_days + 6) // 7

custom_rest, custom_errors = parse_date_lines(custom_rest_input)
holidays, holiday_errors = parse_date_lines(holiday_input)

if custom_errors:
    st.warning(
        f"以下排休日期格式无法识别：{custom_errors}。请使用 2026-02-15 或 2026-02-15~2026-02-23。"
    )

if holiday_errors:
    st.warning(
        f"以下节假日格式无法识别：{holiday_errors}。请使用 2026-02-15 或 2026-02-15~2026-02-23。"
    )


def get_week_dates(index):
    week_start = start_date + timedelta(days=index * 7)
    week_end = week_start + timedelta(days=6)

    if week_end > end_date:
        week_end = end_date

    dates = []
    current = week_start

    while current <= week_end:
        dates.append(current)
        current += timedelta(days=1)

    return dates


def get_week_rest_dates(index, days):
    rest = set()

    if rest_mode in ["双休", "单休"]:
        for d in days:
            if d.weekday() in fixed_rest_weekdays:
                rest.add(d)

    elif rest_mode == "大小周":
        is_big_week = (index % 2 == 0) if first_week_big else (index % 2 == 1)
        current_rest_weekdays = big_week_rest if is_big_week else small_week_rest

        for d in days:
            if d.weekday() in current_rest_weekdays:
                rest.add(d)

    elif rest_mode == "排休/自定义日期":
        pass

    elif rest_mode == "无固定休息":
        pass

    for d in days:
        if d in custom_rest or d in holidays:
            rest.add(d)

    return rest


all_rest = set()

for i in range(week_count):
    days = get_week_dates(i)
    all_rest.update(get_week_rest_dates(i, days))

work_days = total_days - len(all_rest)
work_hours_total = work_days * work_hours_per_day

st.success(
    f"总天数：{total_days}天｜周数：{week_count}周｜休息：{len(all_rest)}天｜"
    f"到岗：{work_days}天｜总工时：{work_hours_total:g}小时"
)

if all_rest:
    st.caption(f"已识别休息日期：{format_date_ranges(all_rest)}")


# =====================
# 自动生成内容
# =====================

def get_auto_work(index, total, zero_work):
    if zero_work:
        return "本周为休息或暂停实习阶段，主要对前期实习内容进行整理和回顾。"

    if index == 0:
        return "熟悉实习单位环境，了解岗位基本要求，学习单位日常工作流程，并逐步适应实习节奏。"

    if index == total - 1:
        return "完成实习收尾工作，对前期实习内容进行整理和总结，并继续配合完成单位安排的相关任务。"

    if index < total / 3:
        return "围绕岗位基础工作展开，逐步熟悉工作流程，并在实际工作中学习相关规范。"

    if index < total * 2 / 3:
        return "在前期熟悉岗位的基础上，继续参与单位日常工作，提升工作熟练度和配合能力。"

    return "继续完成实习岗位相关工作，并对前期工作经验进行总结，进一步提升岗位适应能力。"


def get_reflection(style):
    if style == "通用正式版":
        return "通过本周实习，我对实际岗位工作有了进一步认识，也体会到实际工作中需要具备责任意识、沟通能力和执行能力。后续我会继续保持认真态度，按要求完成实习任务。"

    if style == "学生实习版":
        return "通过本周实习，我认识到学校学习和实际岗位之间存在一定差别。实际工作不仅需要掌握基础知识，也需要认真细致的态度和良好的沟通能力。"

    return "本周实习让我进一步熟悉了岗位工作内容，也提升了自己的实践能力和岗位适应能力。"


def build_records(progress_bar=None, status_text=None):
    input_lines = [line.strip() for line in weekly_input.splitlines() if line.strip()]
    reflection = get_reflection(style)

    records = []

    for i in range(week_count):
        days = get_week_dates(i)

        if not days:
            continue

        rest_dates = get_week_rest_dates(i, days)
        week_work_days = len(days) - len(rest_dates)
        week_work_hours = week_work_days * work_hours_per_day

        if i < len(input_lines):
            work = input_lines[i]
        else:
            work = get_auto_work(i, week_count, week_work_days == 0)

        records.append({
            "周次": f"第{i + 1}周",
            "时间": f"{format_date(days[0])}—{format_date(days[-1])}",
            "休息时间": format_date_ranges(rest_dates),
            "到岗天数": f"{week_work_days}天",
            "工作小时": f"{week_work_hours:g}小时",
            "主要工作内容": work,
            "实习收获与体会": reflection
        })

        if progress_bar is not None:
            progress = int(((i + 1) / week_count) * 100)
            progress_bar.progress(progress)

        if status_text is not None:
            status_text.info(f"正在生成第 {i + 1} 周 / 共 {week_count} 周...")

        time.sleep(0.04)

    if status_text is not None:
        status_text.success("实习周记生成完成！")

    return records


# =====================
# Word 导出格式
# =====================

def set_font(run, name="宋体", size=8.5, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))

    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)

    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_row_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_cell_text(cell, text, bold=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = Pt(14)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    run = p.add_run(str(text))
    set_font(run, bold=bold)


def create_word(records):
    doc = Document()

    # 横向 A4
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)

    # 页边距
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    # 默认字体
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(8.5)

    # 标题
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("岗位实习周记表")
    set_font(title_run, name="黑体", size=16, bold=True)

    # 基本信息
    info_lines = [
        f"实习单位：{unit}",
        f"实习岗位：{job}",
        f"实习时间：{format_date(start_date)}—{format_date(end_date)}",
        f"休息制度：{rest_mode}；每日工时：{work_hours_per_day:g}小时",
        f"总天数：{total_days}天；周数：{week_count}周；休息：{len(all_rest)}天；到岗：{work_days}天；总工时：{work_hours_total:g}小时"
    ]

    for line in info_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        set_font(run, size=10)

    # 表格
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = [
        "周次",
        "时间",
        "休息时间",
        "到岗天数",
        "工作小时",
        "主要工作内容",
        "实习收获与体会"
    ]

    widths = [1.4, 3.2, 4.0, 1.7, 1.8, 7.5, 7.5]

    header_cells = table.rows[0].cells
    set_row_no_split(table.rows[0])

    for i, header in enumerate(headers):
        set_cell_width(header_cells[i], widths[i])
        add_cell_text(header_cells[i], header, bold=True)

    for record in records:
        row = table.add_row()
        set_row_no_split(row)

        cells = row.cells
        values = [
            record["周次"],
            record["时间"],
            record["休息时间"],
            record["到岗天数"],
            record["工作小时"],
            record["主要工作内容"],
            record["实习收获与体会"]
        ]

        for i, value in enumerate(values):
            set_cell_width(cells[i], widths[i])
            add_cell_text(cells[i], value)

    file = BytesIO()
    doc.save(file)
    file.seek(0)
    return file


# =====================
# 生成和下载
# =====================

if "records" not in st.session_state:
    st.session_state.records = None

if "word_file" not in st.session_state:
    st.session_state.word_file = None


if st.button("✅ 一键生成实习周记", type="primary"):
    if not unit.strip():
        st.warning("请先填写实习单位。")
        st.stop()

    if not job.strip():
        st.warning("请先填写实习岗位。")
        st.stop()

    st.markdown("### 生成进度")
    progress_bar = st.progress(0)
    status_text = st.empty()

    records = build_records(progress_bar, status_text)

    st.session_state.records = records
    st.session_state.word_file = create_word(records)

    st.success("生成完成，可以预览或下载 Word 文件。")


if st.session_state.records:
    st.subheader("📄 生成结果预览")

    for record in st.session_state.records:
        with st.expander(f"{record['周次']}｜{record['时间']}"):
            st.write(f"🕒 休息时间：{record['休息时间']}")
            st.write(f"📅 到岗天数：{record['到岗天数']}")
            st.write(f"⏱️ 工作小时：{record['工作小时']}")
            st.write(f"✍️ 主要工作内容：{record['主要工作内容']}")
            st.write(f"📌 实习收获与体会：{record['实习收获与体会']}")

    st.download_button(
        label="📥 下载 Word 表格版周记",
        data=st.session_state.word_file,
        file_name="岗位实习周记表.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


st.markdown("---")
st.caption(
    "本工具用于辅助生成岗位实习周记，适合学生实习记录整理、实习材料归档和 Word 表格周记制作。生成内容仅作参考，提交前建议结合个人真实实习情况进行修改。"
)